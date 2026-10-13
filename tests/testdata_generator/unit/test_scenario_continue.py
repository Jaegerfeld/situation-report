# =============================================================================
# Autor:          Robert Seebauer
# Repository:     https://github.com/Jaegerfeld/situation-report
# KI-Unterstützung: Erstellt mit Unterstützung von Claude (Anthropic)
# Erstellt:       13.10.2026
# Geändert:       13.10.2026
# Lizenz:         BSD-3-Clause (siehe LICENSE)
#
# Fachliche Funktion:
#   Tests der Szenario-Fortschreibung (Schritt 3): Datenfenster der Configs
#   wandern mit gleicher Länge, jeder Roh-Abzug trägt die Kette, die
#   veraltete Quelle bleibt genauso veraltet, Register bleiben byteidentisch,
#   die Snapshots rotieren (altes now = neues prev) und das Delta-Briefing
#   zwischen beiden zeigt vergangene Zeit, transform_data-Ausgaben wachsen,
#   PI-Config reicht bis zum neuen Referenzdatum, Fehlerpfade.
# =============================================================================

from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from portfolio.delta import compute_delta
from portfolio.snapshot import load_snapshot
from portfolio.solution_config import load_solution_config
from testdata_generator.continuation import read_meta
from testdata_generator.scenario import build_portfolio_scenario
from testdata_generator.scenario_continue import (
    ScenarioContinuationError,
    continue_portfolio_scenario,
)

REF = date(2025, 6, 30)
NEW = date(2025, 8, 15)


def _rows(xlsx: Path) -> int:
    return load_workbook(xlsx, read_only=True).active.max_row


@pytest.fixture(scope="module")
def room(tmp_path_factory):
    out = tmp_path_factory.mktemp("room")
    build_portfolio_scenario(out, seed=42, reference=REF, log=lambda m: None)
    before = {
        "registers": {p: p.read_bytes() for p in out.glob("solutions/*/registers/*.json")},
        "snapshot_now": (out / "snapshots" / "snapshot_now.json").read_bytes(),
        "rows": {p: _rows(p) for p in out.glob("solutions/*/arts/*_IssueTimes.xlsx")},
        "raw": {p: json.loads(p.read_text(encoding="utf-8")) for p in out.glob("raw/*_jira.json")},
    }
    lines: list[str] = []
    paths = continue_portfolio_scenario(out, NEW, log=lines.append)
    return out, before, paths, lines


class TestWindowsAndRecords:
    def test_configs_keep_the_window_length_and_move_to_the_new_date(self, room) -> None:
        out, _, _, _ = room
        for cfg_path in (out / "portfolio.json", out / "solutions" / "alpha" / "solution.json",
                         out / "solutions" / "beta" / "solution.json"):
            cfg = load_solution_config(cfg_path)
            assert cfg.to_date == NEW
            assert (cfg.to_date - cfg.from_date).days == 365

    def test_every_raw_extract_carries_the_chain(self, room) -> None:
        out, before, _, _ = room
        for raw, old in before["raw"].items():
            new = json.loads(raw.read_text(encoding="utf-8"))
            meta = read_meta(new)
            assert meta["continued_from"] == [read_meta(old)["to_date"]]
            assert len(new["issues"]) >= len(old["issues"])
            # the past of every old issue is still there
            old_by = {i["key"]: i for i in old["issues"]}
            for issue in new["issues"][:len(old["issues"])]:
                n = len(old_by[issue["key"]]["changelog"]["histories"])
                assert issue["changelog"]["histories"][:n] == old_by[issue["key"]]["changelog"]["histories"]

    def test_stale_source_stays_exactly_as_stale(self, room) -> None:
        out, before, _, _ = room
        beta3 = json.loads((out / "raw" / "Beta-3_jira.json").read_text(encoding="utf-8"))
        assert read_meta(beta3)["to_date"] == (NEW - (REF - date.fromisoformat(
            read_meta(before["raw"][out / "raw" / "Beta-3_jira.json"])["to_date"]))).isoformat()
        alpha1 = json.loads((out / "raw" / "Alpha-1_jira.json").read_text(encoding="utf-8"))
        assert read_meta(alpha1)["to_date"] == NEW.isoformat()

    def test_pi_config_reaches_the_new_reference(self, room) -> None:
        out, _, _, _ = room
        intervals = json.loads((out / "pi_config.json").read_text(encoding="utf-8"))["intervals"]
        assert intervals[-1]["to"] == NEW.isoformat()
        assert intervals[0]["from"] == (NEW - timedelta(days=365)).isoformat()


class TestUntouchedAndRotated:
    def test_registers_are_byte_identical(self, room) -> None:
        out, before, _, _ = room
        for path, raw in before["registers"].items():
            assert path.read_bytes() == raw

    def test_snapshots_rotate(self, room) -> None:
        out, before, paths, _ = room
        assert paths["snapshot_prev"].read_bytes() == before["snapshot_now"]
        now = load_snapshot(paths["snapshot_now"])
        prev = load_snapshot(paths["snapshot_prev"])
        assert now.as_of == NEW and prev.as_of == REF

    def test_delta_between_the_rotated_snapshots_shows_elapsed_time(self, room) -> None:
        _, _, paths, _ = room
        delta = compute_delta(load_snapshot(paths["snapshot_prev"]),
                              load_snapshot(paths["snapshot_now"]))
        assert delta.period_days == (NEW - REF).days
        assert delta.completed_delta > 0
        assert not delta.quiet


class TestOutputsAndLog:
    def test_issue_times_grow(self, room) -> None:
        out, before, _, _ = room
        for path, rows in before["rows"].items():
            assert _rows(path) >= rows

    def test_readme_and_log_tell_the_story(self, room) -> None:
        out, _, _, lines = room
        assert "Fortgeschrieben am" in (out / "README.md").read_text(encoding="utf-8")
        assert any("Snapshots rotiert" in li for li in lines)
        assert any("Beta-3" in li and "Datenstand -60d bleibt" in li for li in lines)


class TestErrors:
    def test_needs_a_later_date(self, room) -> None:
        out, _, _, _ = room
        with pytest.raises(ScenarioContinuationError, match="must lie after"):
            continue_portfolio_scenario(out, NEW, log=lambda m: None)

    def test_needs_a_portfolio(self, tmp_path) -> None:
        with pytest.raises(ScenarioContinuationError, match="portfolio.json"):
            continue_portfolio_scenario(tmp_path, NEW, log=lambda m: None)

    def test_chain_twice(self, room) -> None:
        out, _, _, _ = room
        later = date(2025, 10, 1)
        paths = continue_portfolio_scenario(out, later, log=lambda m: None)
        assert load_snapshot(paths["snapshot_prev"]).as_of == NEW
        assert load_snapshot(paths["snapshot_now"]).as_of == later
        alpha1 = json.loads((out / "raw" / "Alpha-1_jira.json").read_text(encoding="utf-8"))
        assert read_meta(alpha1)["continued_from"] == [REF.isoformat(), NEW.isoformat()]
