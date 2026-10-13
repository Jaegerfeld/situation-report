# =============================================================================
# Autor:          Robert Seebauer
# Repository:     https://github.com/Jaegerfeld/situation-report
# KI-Unterstützung: Erstellt mit Unterstützung von Claude (Anthropic)
# Erstellt:       13.10.2026
# Geändert:       13.10.2026
# Lizenz:         BSD-3-Clause (siehe LICENSE)
#
# Fachliche Funktion:
#   Tests des Lernens aus echten Transitions (Schritt 4): Profil aus dem
#   echten ART_A (Startverteilung, Matrix mit Endzustand, Verweildauern),
#   JSON-Roundtrip ohne Schlüssel, CSV-Lesen, Passung zum Workflow,
#   Pfad-Sampling, und — der eigentliche Zweck — ein mit dem Profil
#   erzeugter Datensatz trägt die Struktur des Originals.
# =============================================================================

from __future__ import annotations

import collections
import json
import random
import statistics
from datetime import date
from pathlib import Path

import pytest

from testdata_generator.generator import GeneratorConfig, generate
from testdata_generator.learn import (
    END,
    LearnError,
    check_profile_fits,
    learn_profile,
    load_profile,
    sample_dwell_hours,
    sample_profile_path,
    save_profile,
)
from testdata_generator.workflow_parser import parse_workflow

ROOT = Path(__file__).resolve().parents[3]
XLSX = ROOT / "tests" / "testdata" / "ART_A" / "ART_A_Transitions.xlsx"
WORKFLOW = ROOT / "testdata_generator" / "workflow_ART_A.txt"


@pytest.fixture(scope="module")
def wf():
    return parse_workflow(WORKFLOW)


@pytest.fixture(scope="module")
def profile(wf):
    return learn_profile(XLSX, stages=wf.stages)


class TestLearn:
    def test_counts_match_the_real_art(self, profile) -> None:
        assert profile.issues == 458
        assert profile.transitions == 1404               # consecutive pairs (rows minus 'Created')
        assert profile.start["Analysis"] == 366
        assert profile.matrix["Analysis"]["Program Backlog"] == 290
        assert profile.matrix["Implementation"]["Implementation"] == 44   # self-loops survive
        assert profile.matrix["Releasing"][END] > 0 and profile.matrix["Done"][END] > 0

    def test_dwell_samples_are_hours_per_source_stage(self, profile) -> None:
        assert profile.dwell_hours["Analysis"] and all(h >= 0 for h in profile.dwell_hours["Analysis"])
        assert len(profile.dwell_hours["Analysis"]) <= 400

    def test_csv_reads_like_xlsx(self, wf, profile, tmp_path) -> None:
        # The CSV twin is built from the tracked XLSX: the module's own
        # ART_A_*.csv samples are gitignored and absent in CI.
        from openpyxl import load_workbook
        ws = load_workbook(XLSX, read_only=True).active
        csv_path = tmp_path / "ART_A_Transitions.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as f:
            for row in ws.iter_rows(values_only=True):
                f.write(";".join("" if c is None else str(c) for c in row) + "\n")
        from_csv = learn_profile(csv_path, stages=wf.stages)
        assert from_csv.to_dict() == {**profile.to_dict(), "source": csv_path.name}

    def test_without_stages_the_seen_order_is_kept(self) -> None:
        p = learn_profile(XLSX)
        assert p.stages[0] in ("Analysis", "Funnel") and "Done" in p.stages

    def test_unknown_stages_are_dropped(self) -> None:
        p = learn_profile(XLSX, stages=["Analysis", "Program Backlog", "Implementation"])
        assert set(p.stages) == {"Analysis", "Program Backlog", "Implementation"}
        assert all(to in ("Analysis", "Program Backlog", "Implementation", END)
                   for row in p.matrix.values() for to in row)

    def test_errors(self, tmp_path) -> None:
        with pytest.raises(LearnError, match="not found"):
            learn_profile(tmp_path / "nope.xlsx")
        empty = tmp_path / "t.csv"
        empty.write_text("Key;Transition;Timestamp\n", encoding="utf-8")
        with pytest.raises(LearnError, match="no transition rows"):
            learn_profile(empty)
        with pytest.raises(LearnError, match="no issue with a known"):
            learn_profile(XLSX, stages=["Nothing"])


class TestProfileFile:
    def test_roundtrip_without_keys_or_dates(self, profile, tmp_path) -> None:
        target = tmp_path / "profile.json"
        save_profile(target, profile)
        text = target.read_text(encoding="utf-8")
        assert "ART_A-" not in text and "2025" not in text.replace("schema", "")
        loaded = load_profile(target)
        assert loaded.to_dict() == profile.to_dict()

    def test_schema_guard(self, tmp_path) -> None:
        bad = tmp_path / "bad.json"
        bad.write_text(json.dumps({"schema": 9}), encoding="utf-8")
        with pytest.raises(LearnError, match="schema"):
            load_profile(bad)

    def test_fit_check(self, profile) -> None:
        check_profile_fits(profile, profile.stages)
        with pytest.raises(LearnError, match="not in the workflow"):
            check_profile_fits(profile, ["Only", "These"])


class TestSampling:
    def test_paths_follow_the_matrix(self, profile) -> None:
        rng = random.Random(1)
        paths = [sample_profile_path(profile, rng) for _ in range(500)]
        assert all(p for p in paths)
        firsts = collections.Counter(p[0] for p in paths)
        assert firsts["Analysis"] / 500 > 0.6
        for p in paths:
            for a, b in zip(p, p[1:]):
                assert profile.matrix[a].get(b, 0) > 0
        assert statistics.median(len(p) for p in paths) <= 6

    def test_dwell_resample_and_fallback(self, profile) -> None:
        rng = random.Random(2)
        assert sample_dwell_hours(profile, "Analysis", rng, 5.0) >= 0.25
        assert sample_dwell_hours(profile, "Nowhere", rng, 5.0) == 5.0


class TestGenerateWithProfile:
    def _seqs(self, data):
        return [[h["items"][0]["toString"] for h in i["changelog"]["histories"]]
                for i in data["issues"]]

    def test_generated_set_carries_the_real_structure(self, wf, profile, tmp_path) -> None:
        target = tmp_path / "profile.json"
        save_profile(target, profile)
        cfg = GeneratorConfig(project_key="L", issue_count=500, from_date=date(2025, 1, 1),
                              to_date=date(2025, 12, 31), seed=4, transition_profile=str(target))
        seqs = self._seqs(generate(wf, cfg))
        idx = {s: i for i, s in enumerate(wf.stages)}
        pairs: collections.Counter = collections.Counter()
        back = loops = 0
        per = []
        for st in seqs:
            per.append(len(st))
            for a, b in zip(st, st[1:]):
                pairs[(a, b)] += 1
                loops += a == b
                back += idx[b] < idx[a]
        tot = sum(pairs.values())
        assert len(pairs) >= 28
        assert 0.07 <= back / tot <= 0.17
        assert 0.04 <= loops / tot <= 0.11
        assert statistics.median(per) <= 5
        # only stages of the workflow, status consistent
        for issue in generate(wf, cfg)["issues"][:50]:
            hist = issue["changelog"]["histories"]
            if hist:
                assert issue["fields"]["status"]["name"] == hist[-1]["items"][0]["toString"]

    def test_profile_with_ct_budget_keeps_the_pattern_machinery(self, wf, profile, tmp_path) -> None:
        target = tmp_path / "profile.json"
        save_profile(target, profile)
        cfg = GeneratorConfig(project_key="L", issue_count=200, from_date=date(2025, 1, 1),
                              to_date=date(2025, 12, 31), seed=4, transition_profile=str(target),
                              mean_cycle_days=20, pattern="cluster")
        data = generate(wf, cfg)
        assert len(data["issues"]) == 200
        assert any(i["changelog"]["histories"] for i in data["issues"])

    def test_missing_or_unfitting_profile_is_a_clean_error(self, wf, tmp_path) -> None:
        cfg = GeneratorConfig(project_key="L", issue_count=5, transition_profile=str(tmp_path / "x.json"))
        with pytest.raises(LearnError):
            generate(wf, cfg)
        other = learn_profile(XLSX)
        other.stages.append("Alien")
        other.start["Alien"] = 1
        target = tmp_path / "alien.json"
        save_profile(target, other)
        with pytest.raises(LearnError, match="not in the workflow"):
            generate(wf, GeneratorConfig(project_key="L", issue_count=5,
                                         transition_profile=str(target)))
