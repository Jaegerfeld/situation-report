# =============================================================================
# Autor:          Robert Seebauer
# Repository:     https://github.com/Jaegerfeld/situation-report
# KI-Unterstützung: Erstellt mit Unterstützung von Claude (Anthropic)
# Erstellt:       26.10.2026
# Geändert:       26.10.2026
# Lizenz:         BSD-3-Clause (siehe LICENSE)
#
# Fachliche Funktion:
#   Unit-Tests für das Portfolio-Szenario mit konfigurierbarer Struktur
#   (Beta Schnitt 1a): die Standard-Vorlage liefert dieselben Daten wie der
#   Aufruf ohne Struktur; kleine Struktur (eine Solution), drei Solutions
#   mit 3/4/1 ARTs (Rollen wandern mit, dritte Solution ohne Anker),
#   mehrere Portfolios in Unterordnern; structure.json/story.json im
#   Datenraum; Roundtrip durch die echten Parser.
# =============================================================================

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from portfolio.aggregator import build_pooled_report_data
from portfolio.dependency_config import DEP_BLOCKED, load_dependencies
from portfolio.risks_config import load_risks
from portfolio.solution_config import KIND_PORTFOLIO, load_solution_config
from testdata_generator.scenario import build_portfolio_scenario
from testdata_generator.story import ROLE_OUTLIER_ART, ROLE_WEAK_SOURCE_ART, load_story
from testdata_generator.structure import (
    REGISTER_NAMES,
    TEMPLATE_SMALL,
    TEMPLATE_STANDARD,
    PortfolioSpec,
    SolutionSpec,
    Structure,
    load_structure,
    structure_from_arts_spec,
    structure_template,
)

REF = date(2025, 6, 30)


def _sheet_values(path):
    ws = load_workbook(path, read_only=True).active
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


class TestStandardTemplateEqualsDefault:
    def test_same_issue_times_and_registers(self, tmp_path) -> None:
        a = tmp_path / "default"
        b = tmp_path / "standard"
        build_portfolio_scenario(a, seed=42, reference=REF, log=lambda m: None,
                                 scale="s")
        build_portfolio_scenario(b, seed=42, reference=REF, log=lambda m: None,
                                 scale="s", structure=structure_template(TEMPLATE_STANDARD))
        for art in ("Alpha-1", "Beta-3"):
            key = "alpha" if art.startswith("Alpha") else "beta"
            assert _sheet_values(a / "solutions" / key / "arts" / f"{art}_IssueTimes.xlsx") \
                == _sheet_values(b / "solutions" / key / "arts" / f"{art}_IssueTimes.xlsx")
        for key in ("alpha", "beta"):
            for reg in REGISTER_NAMES:
                pa = a / "solutions" / key / "registers" / f"{reg}.json"
                pb = b / "solutions" / key / "registers" / f"{reg}.json"
                assert pa.read_text(encoding="utf-8") == pb.read_text(encoding="utf-8"), reg

    def test_data_room_carries_structure_and_story(self, tmp_path) -> None:
        paths = build_portfolio_scenario(tmp_path, seed=42, reference=REF,
                                         log=lambda m: None, scale="s")
        structure = load_structure(paths["structure"])
        assert structure.describe() == "1 Portfolio · 2 Solutions · 6 ARTs"
        story = load_story(paths["story"])
        assert story.roles[ROLE_OUTLIER_ART] == "Alpha-3"
        assert story.roles[ROLE_WEAK_SOURCE_ART] == "Beta-3"
        assert len(story.events) == 7
        readme = paths["readme"].read_text(encoding="utf-8")
        assert "## Struktur" in readme and "Standard-Lage" in readme
        assert "ART Alpha-3 ist der Ausreißer" in readme
        # Generic and legacy keys both point at the same files.
        assert paths["flow_problems_alpha"] == paths["flow_alpha"]
        assert paths["solutions"]["beta"] == paths["solution_beta"]


@pytest.fixture(scope="module")
def small(tmp_path_factory):
    out = tmp_path_factory.mktemp("small")
    messages: list[str] = []
    paths = build_portfolio_scenario(
        out, seed=7, reference=REF, log=messages.append, scale="s",
        structure=structure_template(TEMPLATE_SMALL))
    return out, paths, messages


@pytest.fixture(scope="module")
def room(tmp_path_factory):
    out = tmp_path_factory.mktemp("three")
    paths = build_portfolio_scenario(
        out, seed=11, reference=REF, log=lambda m: None, scale="s",
        structure=structure_from_arts_spec("3,4s,1"))
    return out, paths


class TestSmallStructure:
    def test_one_solution_two_arts(self, small) -> None:
        out, paths, _ = small
        assert sorted(p.name for p in (out / "raw").glob("*_jira.json")) \
            == ["Alpha-1_jira.json", "Alpha-2_jira.json"]
        assert not (out / "solutions" / "beta").exists()
        portfolio = load_solution_config(paths["portfolio"])
        assert portfolio.kind == KIND_PORTFOLIO and len(portfolio.members) == 1

    def test_outlier_moves_to_the_last_art(self, small) -> None:
        _, paths, _ = small
        story = load_story(paths["story"])
        assert story.roles[ROLE_OUTLIER_ART] == "Alpha-2"
        deps = load_dependencies(paths["dependencies_alpha"])
        blocked = [d for d in deps.dependencies if d.status == DEP_BLOCKED]
        assert [(d.from_art, d.to_art) for d in blocked] == [("ART Alpha-1", "ART Alpha-2")]
        # No cross-solution seam without a second solution.
        assert not any(d.dep_id == "AD-4" for d in deps.dependencies)

    def test_notice_is_logged_and_in_story(self, small) -> None:
        _, paths, messages = small
        assert any("Nur eine Solution" in m for m in messages)
        assert any("Nur eine Solution" in n for n in load_story(paths["story"]).notices)

    def test_snapshots_and_pooled_report_work(self, small) -> None:
        _, paths, _ = small
        assert paths["snapshot_prev"].is_file() and paths["snapshot_now"].is_file()
        cfg = load_solution_config(paths["portfolio"])
        data = build_pooled_report_data(cfg, log=lambda m: None)
        assert data is not None


class TestThreeSolutions:
    def test_art_counts_follow_the_structure(self, room) -> None:
        out, paths = room
        assert len(list((out / "raw").glob("*_jira.json"))) == 8
        assert len(list(out.glob("solutions/*/arts/*_IssueTimes.xlsx"))) == 8
        assert sorted(paths["solutions"]) == ["alpha", "beta", "gamma"]
        assert len(load_solution_config(paths["solution_beta"]).members) == 4
        assert len(load_solution_config(paths["solution_gamma"]).members) == 1

    def test_workflow_per_solution(self, room) -> None:
        _, paths = room
        assert load_solution_config(paths["solution_alpha"]).stage_map is None
        assert load_solution_config(paths["solution_beta"]).stage_map is not None
        assert load_solution_config(paths["solution_gamma"]).stage_map is None

    def test_roles_bind_to_positions(self, room) -> None:
        _, paths = room
        story = load_story(paths["story"])
        assert story.roles[ROLE_OUTLIER_ART] == "Alpha-3"
        assert story.roles[ROLE_WEAK_SOURCE_ART] == "Beta-4"
        beta_risks = load_risks(paths["risks_beta"])
        assert any("ART Beta-4" in r.title for r in beta_risks.risks)
        assert any("Solution Gamma" in n for n in story.notices)

    def test_third_solution_has_registers_without_anchors(self, room) -> None:
        _, paths = room
        for reg in REGISTER_NAMES:
            assert paths[f"{reg}_gamma"].is_file()
        # Scale s = anchors only, and Gamma has none -> empty risk register.
        assert load_risks(paths["risks_gamma"]).risks == []
        # A one-ART solution has no inner dependencies.
        assert load_dependencies(paths["dependencies_gamma"]).dependencies == []

    def test_portfolio_report_pools_all_three(self, room) -> None:
        _, paths = room
        cfg = load_solution_config(paths["portfolio"])
        assert [m.name for m in cfg.members] == ["Solution Alpha", "Solution Beta",
                                                 "Solution Gamma"]
        assert build_pooled_report_data(cfg, log=lambda m: None) is not None


class TestSeveralPortfolios:
    def test_each_portfolio_is_its_own_data_room(self, tmp_path) -> None:
        structure = Structure([
            PortfolioSpec("Portfolio Nord", [SolutionSpec("Alpha", ["Alpha-1", "Alpha-2"])]),
            PortfolioSpec("Portfolio Süd", [SolutionSpec("Beta", ["Beta-1"]),
                                            SolutionSpec("Gamma", ["Gamma-1"])]),
        ])
        paths = build_portfolio_scenario(tmp_path, seed=3, reference=REF,
                                         log=lambda m: None, scale="s",
                                         structure=structure)
        assert sorted(paths["portfolios"]) == ["portfolio_nord", "portfolio_sued"]
        nord = paths["portfolios"]["portfolio_nord"]
        assert nord["portfolio"] == tmp_path / "portfolio_nord" / "portfolio.json"
        assert load_solution_config(nord["portfolio"]).name == "Portfolio Nord"
        sued = paths["portfolios"]["portfolio_sued"]
        assert (tmp_path / "portfolio_sued" / "structure.json").is_file()
        assert load_structure(sued["structure"]).n_portfolios == 1
        assert load_structure(paths["structure"]).n_portfolios == 2
        assert not (tmp_path / "portfolio.json").exists()
        readme = paths["readme"].read_text(encoding="utf-8")
        assert "portfolio_nord/" in readme and "Portfolio Süd" in readme
        # Both stands exist per portfolio.
        assert sued["snapshot_now"].is_file() and nord["snapshot_prev"].is_file()


class TestInvalidStructure:
    def test_build_rejects_invalid_structure(self, tmp_path) -> None:
        bad = Structure([PortfolioSpec("P", [SolutionSpec("A", [])])])
        with pytest.raises(ValueError, match="at least one ART"):
            build_portfolio_scenario(tmp_path, reference=REF, log=lambda m: None,
                                     structure=bad)
