# =============================================================================
# Autor:          Robert Seebauer
# Repository:     https://github.com/Jaegerfeld/situation-report
# KI-Unterstützung: Erstellt mit Unterstützung von Claude (Anthropic)
# Erstellt:       13.10.2026
# Geändert:       13.10.2026
# Lizenz:         BSD-3-Clause (siehe LICENSE)
#
# Fachliche Funktion:
#   Lernen aus echten Daten (Schritt 4 der Generator-Überarbeitung): Aus
#   einer Transitions-Datei (XLSX oder CSV aus transform_data: Key,
#   Transition, Timestamp) wird ein Übergangsprofil geschätzt — Verteilung
#   der ersten Stufe, Übergangsmatrix erster Ordnung mit einem Endzustand
#   (wo Vorgänge stehen bleiben), Verweildauern je Ausgangsstufe. Das Profil
#   enthält nur Stufennamen, Zählungen und Stunden — keine Schlüssel, keine
#   Texte, keine Zeitpunkte: Was damit erzeugt wird, teilt mit der Quelle
#   nur die Struktur. Der Generator würfelt mit dem Profil Pfad und
#   Verweildauern statt mit den Reglern; Fertig- und To-Do-Quote ergeben
#   sich aus den Daten.
# =============================================================================

from __future__ import annotations

import csv
import json
import random
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

PROFILE_SCHEMA = 1
END = "__end__"
CREATED = "Created"
#: Verweildauer-Stichproben je Stufe, die das Profil höchstens behält.
MAX_DWELL_SAMPLES = 400


class LearnError(ValueError):
    """The transitions file cannot be read or yields no usable transitions."""


@dataclass
class TransitionProfile:
    """What the generator needs from a real ART — structure only."""
    source: str
    issues: int
    stages: list[str]
    start: dict[str, int]                        # first target after Created
    matrix: dict[str, dict[str, int]]            # from -> to (or END) -> count
    dwell_hours: dict[str, list[float]]          # from -> observed hours
    schema: int = PROFILE_SCHEMA
    note: str = field(default="counts and hours only - no keys, no texts, no dates")

    @property
    def transitions(self) -> int:
        return sum(n for row in self.matrix.values() for to, n in row.items() if to != END)

    def to_dict(self) -> dict[str, Any]:
        return {"schema": self.schema, "source": self.source, "note": self.note,
                "issues": self.issues, "stages": self.stages, "start": self.start,
                "matrix": self.matrix, "dwell_hours": self.dwell_hours}


def _parse_ts(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise LearnError(f"timestamp not understood: {value!r}")


def _read_rows(path: Path) -> list[tuple[str, str, datetime]]:
    """(key, transition, timestamp) rows from an XLSX or CSV transitions file."""
    path = Path(path)
    if not path.is_file():
        raise LearnError(f"transitions file not found: {path}")
    rows: list[tuple[str, str, datetime]] = []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True).active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or row[0] is None:
                continue
            rows.append((str(row[0]), str(row[1]), _parse_ts(row[2])))
    else:
        with path.open(encoding="utf-8-sig", newline="") as f:
            sample = f.read(2048)
            f.seek(0)
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
            reader = csv.reader(f, delimiter=delimiter)
            for i, row in enumerate(reader):
                if i == 0 or len(row) < 3 or not row[0]:
                    continue
                rows.append((row[0], row[1], _parse_ts(row[2])))
    if not rows:
        raise LearnError(f"no transition rows in {path.name}")
    return rows


def learn_profile(path: Path, stages: list[str] | None = None) -> TransitionProfile:
    """
    Estimate the profile from a transitions file.

    With ``stages`` (the workflow's canonical names) transitions into unknown
    stages are dropped, so the profile fits that workflow; without, the
    stages are those seen, in order of first appearance.
    """
    rows = _read_rows(Path(path))
    by_key: dict[str, list[tuple[str, datetime]]] = defaultdict(list)
    for key, label, ts in rows:
        by_key[key].append((label, ts))
    known = set(stages) if stages else None
    seen: list[str] = list(stages) if stages else []
    start: Counter = Counter()
    matrix: dict[str, Counter] = defaultdict(Counter)
    dwell: dict[str, list[float]] = defaultdict(list)
    issues = 0
    for seq in by_key.values():
        seq.sort(key=lambda x: x[1])
        steps = [(label, ts) for label, ts in seq
                 if label != CREATED and (known is None or label in known)]
        if not steps:
            continue
        issues += 1
        for label, _ in steps:
            if label not in seen:
                seen.append(label)
        start[steps[0][0]] += 1
        for (a, ta), (b, tb) in zip(steps, steps[1:]):
            matrix[a][b] += 1
            hours = max(0.0, (tb - ta).total_seconds() / 3600.0)
            if len(dwell[a]) < MAX_DWELL_SAMPLES:
                dwell[a].append(round(hours, 2))
        matrix[steps[-1][0]][END] += 1
    if issues == 0:
        raise LearnError("no issue with a known transition - check the workflow stages")
    return TransitionProfile(
        source=Path(path).name, issues=issues, stages=seen,
        start=dict(start), matrix={k: dict(v) for k, v in matrix.items()},
        dwell_hours=dict(dwell))


def save_profile(path: Path, profile: TransitionProfile) -> None:
    Path(path).write_text(json.dumps(profile.to_dict(), indent=2, ensure_ascii=False),
                          encoding="utf-8")


def load_profile(path: Path) -> TransitionProfile:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise LearnError(f"cannot read profile {path}: {exc}") from exc
    if data.get("schema") != PROFILE_SCHEMA:
        raise LearnError(f"profile {path}: unsupported schema {data.get('schema')!r}")
    return TransitionProfile(
        source=str(data.get("source", "")), issues=int(data.get("issues", 0)),
        stages=list(data.get("stages", [])), start={k: int(v) for k, v in data["start"].items()},
        matrix={k: {kk: int(vv) for kk, vv in row.items()} for k, row in data["matrix"].items()},
        dwell_hours={k: [float(x) for x in v] for k, v in data.get("dwell_hours", {}).items()},
        note=str(data.get("note", "")))


def check_profile_fits(profile: TransitionProfile, stages: list[str]) -> None:
    """Every stage the profile knows must exist in the workflow."""
    unknown = [s for s in profile.stages if s not in stages]
    if unknown:
        raise LearnError(
            f"profile stages not in the workflow: {', '.join(unknown)} - "
            f"learn the profile with the workflow's stages or adjust the workflow.")


def _weighted(rng: random.Random, counts: dict[str, int]) -> str | None:
    items = [(k, n) for k, n in counts.items() if n > 0]
    if not items:
        return None
    return rng.choices([k for k, _ in items], weights=[n for _, n in items], k=1)[0]


def sample_profile_path(profile: TransitionProfile, rng: random.Random,
                        max_steps: int = 60) -> list[str]:
    """
    The stage names an issue visits, drawn from the profile: first stage from
    ``start``, then the matrix until END (or ``max_steps``). An empty list
    means the issue never left its creation state.
    """
    cur = _weighted(rng, profile.start)
    if cur is None:
        return []
    path = [cur]
    for _ in range(max_steps):
        nxt = _weighted(rng, profile.matrix.get(cur, {}))
        if nxt is None or nxt == END:
            break
        path.append(nxt)
        cur = nxt
    return path


def sample_dwell_hours(profile: TransitionProfile, stage: str, rng: random.Random,
                       fallback: float) -> float:
    """Resample an observed dwell of ``stage`` (hours); ``fallback`` without data."""
    samples = profile.dwell_hours.get(stage)
    if not samples:
        return fallback
    return max(0.25, rng.choice(samples))
